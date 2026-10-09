"""groundtruth/xlsx.py: Excel cho QA duyệt Ground-Truth (YAML là nguồn sự thật; xlsx hai chiều, so theo NGỮ NGHĨA, gộp ba chiều).

Phần 1 test xlsx.py thuần (catalog dựng tay, QA sửa Excel bằng openpyxl). Phần 2 test cổng `gt validate`; phần 3 test CLI (generate/regen/import-xlsx/export-xlsx).
"""
import copy
import io
import json
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from qc_agent.groundtruth import check as gt_check
from qc_agent.groundtruth import coverage as cov
from qc_agent.groundtruth import render as gt_render
from qc_agent.groundtruth import schema as gt_schema
from qc_agent.groundtruth import xlsx
from qc_agent.groundtruth.merge import merge, protected
from tests.test_gt_cli import GT, approve_everything, env, fake, generate_args, generated, gt, load_catalog, make_sut, save_catalog  # noqa: F401
from tests.test_gt_coverage import FACTS

SHA = "a" * 64


def step(method, path, status, *, body=None, path_params=None, capture=None, asserts=None):
    request = {"method": method, "path": path}
    if path_params:
        request["path_params"] = path_params
    if body is not None:
        request["json"] = body
    out = {"request": request, "expect": {"status": list(status)}}
    if asserts:
        out["expect"]["json"] = asserts
    if capture:
        out["capture"] = capture
    return out


def mini() -> dict:
    asserts = [{"path": "$.id", "op": "type", "value": "integer"}, {"path": "$.title", "op": "eq", "value": "  x  "}, {"path": "$.n", "op": "eq", "value": 7},
               {"path": "$.s", "op": "eq", "value": "007"}, {"path": "$.v", "op": "eq", "value": "1.10"}, {"path": "$.z", "op": "eq", "value": None},
               {"path": "$.b", "op": "eq", "value": True}, {"path": "$.t", "op": "eq", "value": "true"}, {"path": "$.l", "op": "len_gte", "value": 3}, {"path": "$.x", "op": "exists"},
               {"path": "$.e", "op": "eq", "value": ""}]
    return {
        "version": 1, "prd": {"id": "demo", "sha256": SHA, "source": "docs/prd/demo.md"}, "generated_by": {"model": "claude-opus-5-5", "prompt_version": "gt-agent/1"},
        "status": "draft",
        "stories": [{"story_id": "US-1", "title": "Ghi chú", "acs": [{"ac_id": "AC-1.1", "text": "tạo"}, {"ac_id": "AC-1.2", "text": "đọc"}]},
                    {"story_id": "US-2", "title": "Liệt kê", "acs": [{"ac_id": "AC-2.1", "text": "danh sách"}]}],
        "test_cases": [
            {"tc_id": "TC-AC-1.1-aaaa01", "title": "tạo ghi chú", "ac_refs": ["AC-1.1"], "kind": "api_functional", "status": "draft", "origin": "llm", "priority": "high",
             "technique": "boundary", "preconditions": "không", "rationale": "biên 200", "evidence": [{"path": "app/main.py", "line": 3}, {"path": "app/m.py"}],
             "steps": [step("POST", "/notes", (201,), body={"title": "  x  ", "body": "b", "n": 1.10, "nested": {"k": [1, "2", None]}}, asserts=asserts)]},
            {"tc_id": "TC-AC-1.2-aaaa02", "title": "tạo rồi đọc", "ac_refs": ["AC-1.2", "AC-1.1"], "kind": "flow", "status": "draft", "origin": "llm",
             "steps": [step("POST", "/notes", (201,), body={"title": "a", "body": "b"}, capture={"note_id": "$.id"}),
                       step("GET", "/notes/{note_id}", (200, 404), path_params={"note_id": "{{note_id}}"})]},
            {"tc_id": "TC-AC-2.1-qa-bbbb03", "title": "danh sách", "ac_refs": ["AC-2.1"], "kind": "api_contract", "status": "approved", "origin": "qa", "notes": "QA tự thêm",
             "steps": [step("GET", "/notes", (200,))]},
        ],
        "uncovered_acs": [],
        "coverage_plan": [{"ac_id": "AC-1.1", "technique": "boundary", "scenario": "title dài 200", "decision": "planned", "tc_ids": ["TC-AC-1.1-aaaa01"]}],
        "waivers": [{"kind": "api", "target": "GET /notes 500", "reason_code": "not_applicable", "reason": "không tái hiện", "status": "draft"}],
        "spec_conflicts": [{"ac_id": "AC-1.2", "summary": "mã khác PRD", "status": "open", "evidence": [{"path": "app/m.py", "line": 9}]}],
    }


def export(catalog=None, facts=None) -> bytes:
    return xlsx.render_xlsx(catalog or mini(), facts)


class Edit:
    """QA sửa Excel bằng openpyxl (như Excel lưu lại): tìm dòng theo giá trị cột, đặt ô, thêm/xoá dòng, xoá sheet."""

    def __init__(self, data: bytes):
        self.wb = load_workbook(io.BytesIO(data))

    def col(self, sheet, name):
        header = [str(c.value).lower() if c.value is not None else "" for c in self.wb[sheet][1]]
        return header.index(name) + 1

    def row(self, sheet, **where) -> int:
        for number in range(2, self.wb[sheet].max_row + 1):
            if all(str(self.wb[sheet].cell(number, self.col(sheet, k)).value) == str(v) for k, v in where.items()):
                return number
        raise KeyError(where)

    def set(self, sheet, where: dict, column: str, value):
        self.wb[sheet].cell(self.row(sheet, **where), self.col(sheet, column)).value = value
        return self

    def add(self, sheet, **values):
        number = self.wb[sheet].max_row + 1
        for name, value in values.items():
            self.wb[sheet].cell(number, self.col(sheet, name)).value = value
        return self

    def delete(self, sheet, **where):
        self.wb[sheet].delete_rows(self.row(sheet, **where))
        return self

    def drop(self, sheet):
        del self.wb[sheet]
        return self

    def bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()


def do_import(catalog, data, *, expect_ok=True):
    theirs, base, syntax = xlsx.read_xlsx(data)
    result = xlsx.import_into(catalog, theirs, base, syntax)
    if expect_ok:
        assert result.ok, (result.conflicts, result.errors)
    return result


# ================= xuất =================

def test_workbook_structure_and_hidden_meta():
    wb = load_workbook(io.BytesIO(export(facts=FACTS)))
    assert wb.sheetnames == ["HuongDan", "Catalog", "TestCases", "Steps", "Assertions", "Uncovered", "Waivers", "SpecConflicts", "Coverage", "AgentPlan", "_meta"]
    assert wb["_meta"].sheet_state == "veryHidden" and all(wb[n].sheet_state == "visible" for n in wb.sheetnames[:-1])
    tc = wb["TestCases"]
    assert [c.value for c in tc[1]] == list(xlsx.TC_COLUMNS) and tc.freeze_panes == "A2"
    assert tc.cell(1, 8).fill.fgColor.rgb.endswith("FFF2CC") and tc.cell(1, 1).fill.fgColor.rgb.endswith("D9D9D9")   # vàng: sửa được; xám: chỉ đọc
    assert {dv.formula1 for dv in tc.data_validations.dataValidation} >= {'"draft,approved,rejected"', '"high,medium,low"', '"api_contract,api_functional,flow"'}
    assert "qc-agent gt import-xlsx" in str(wb["HuongDan"]["A3"].value)
    rows = {r[0].value: [c.value for c in r] for r in tc.iter_rows(min_row=2)}
    assert rows["TC-AC-1.1-aaaa01"][1:4] == ["US-1", "AC-1.1", "tạo ghi chú"] and rows["TC-AC-1.2-aaaa02"][2] == "AC-1.2, AC-1.1"
    assert rows["TC-AC-1.2-aaaa02"][14] == 2 or rows["TC-AC-1.2-aaaa02"][14] == "2"
    assert "POST /notes -> 201 | GET /notes/{note_id} -> 200,404" == rows["TC-AC-1.2-aaaa02"][15]


def test_render_is_byte_identical_across_seconds():
    """openpyxl đặt `modified` và timestamp zip = giờ hiện tại lúc save; DOS time của zip có độ phân giải 2 giây nên phải ngủ qua ranh giới đó."""
    first = export(facts=FACTS)
    time.sleep(2.1)
    assert export(facts=FACTS) == first


def test_zip_members_have_fixed_timestamps_and_core_modified():
    with zipfile.ZipFile(io.BytesIO(export(facts=FACTS))) as archive:
        assert {item.date_time for item in archive.infolist()} == {(2000, 1, 1, 0, 0, 0)}
        core = archive.read("docProps/core.xml")
    assert b"2000-01-01T00:00:00Z</dcterms:modified>" in core and b"2000-01-01T00:00:00" in core.split(b"dcterms:created")[1]


def test_deterministic_xlsx_still_round_trips(tmp_path):
    catalog = mini()
    path = tmp_path / "t.xlsx"
    path.write_bytes(export(catalog, FACTS))
    theirs, base, findings = xlsx.read_xlsx(path)
    assert not findings and xlsx.strip_rows(theirs) == xlsx.projection(catalog) and base == xlsx.projection(catalog)


def test_every_cell_is_text_so_formulas_and_excel_coercion_cannot_happen():
    catalog = mini()
    evil = ['=HYPERLINK("http://evil.example","x")', "+cmd|' /C calc'!A0", "@SUM(1+1)", "-2+3", "=1+1"]
    catalog["test_cases"][0]["title"] = evil[0]
    catalog["test_cases"][0]["notes"] = evil[1]
    catalog["test_cases"][1]["rationale"] = evil[2]
    catalog["test_cases"][1]["preconditions"] = evil[3]
    catalog["test_cases"][2]["notes"] = evil[4]
    data = export(catalog)
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate)
                if cell.value is not None and cell.row > 1 and ws.title != "HuongDan":    # dòng 1 là tiêu đề do code đặt; dữ liệu người dùng từ dòng 2
                    assert cell.number_format == "@", (ws.title, cell.coordinate)
    assert wb["TestCases"]["D2"].value == evil[0] and wb["TestCases"]["J2"].value == evil[1]
    theirs, base, syntax = xlsx.read_xlsx(data)
    assert syntax == [] and xlsx.strip_rows(theirs) == xlsx.projection(catalog)               # chuỗi giống công thức đi vòng nguyên vẹn


def test_text_that_excel_would_mangle_survives_the_round_trip():
    catalog = mini()
    theirs, base, syntax = xlsx.read_xlsx(export(catalog))
    values = {a["path"]: a["value"] for a in xlsx.strip_rows(theirs)["test_cases"]["TC-AC-1.1-aaaa01"]["steps"][0]["expect"]["json"] if "value" in a}
    assert values["$.s"] == "007" and values["$.v"] == "1.10" and values["$.t"] == "true" and values["$.z"] is None and values["$.b"] is True and values["$.n"] == 7
    assert values["$.title"] == "  x  " and values["$.e"] == "" and isinstance(values["$.n"], int)          # khoảng trắng đầu/cuối là NỘI DUNG của assertion
    body = xlsx.strip_rows(theirs)["test_cases"]["TC-AC-1.1-aaaa01"]["steps"][0]["request"]["json"]
    assert body == {"title": "  x  ", "body": "b", "n": 1.1, "nested": {"k": [1, "2", None]}}
    assert syntax == [] and xlsx.strip_rows(theirs) == xlsx.projection(catalog)


def test_unicode_newlines_and_emoji_round_trip():
    catalog = mini()
    catalog["test_cases"][0]["notes"] = "dòng 1\ndòng 2 🚀 日本語"
    theirs, _, _ = xlsx.read_xlsx(export(catalog))
    assert xlsx.strip_rows(theirs)["test_cases"]["TC-AC-1.1-aaaa01"]["notes"] == "dòng 1 dòng 2 🚀 日本語"   # projection co khoảng trắng (cả hai phía như nhau)
    assert xlsx.strip_rows(theirs) == xlsx.projection(catalog)


def test_a_cell_too_large_for_excel_is_an_error_without_the_content():
    catalog = mini()
    catalog["test_cases"][0]["steps"][0]["request"]["json"] = {"blob": "SECRET" * 8000}
    with pytest.raises(xlsx.XlsxError) as error:
        export(catalog)
    assert "SECRET" not in str(error.value) and "32000" in str(error.value)


def test_invalid_catalog_is_refused_before_rendering():
    catalog = mini()
    catalog["test_cases"][0]["kind"] = "ui"
    with pytest.raises(xlsx.XlsxError):
        export(catalog)


def test_coverage_sheet_lists_acs_and_the_technique_and_api_cells(tmp_path):
    catalog = mini()
    catalog["test_cases"][2]["steps"] = [step("GET", "/items", (200,))]   # chạm operation của FACTS
    wb = load_workbook(io.BytesIO(export(catalog, FACTS)))
    rows = [[c.value for c in r] for r in wb["Coverage"].iter_rows(min_row=2)]
    assert [r[:3] for r in rows if r[0] == "ac"] == [["ac", "AC-1.1", "covered"], ["ac", "AC-1.2", "covered"], ["ac", "AC-2.1", "covered"]]
    assert any(r[0] == "api" and r[1] == "GET /items 200" and r[2] == "covered" and r[3] == "TC-AC-2.1-qa-bbbb03" for r in rows)
    assert any(r[0] == "technique" and r[1] == "POST /items boundary:max_length:body.name@11" and r[2] == "gap" for r in rows)
    assert any(r[0] == "api" and r[1] == "POST /items 500" and r[2] == "exempt" for r in rows)
    without = load_workbook(io.BytesIO(export(catalog)))
    assert {r[0].value for r in without["Coverage"].iter_rows(min_row=2)} == {"ac"}
    plan = [[c.value for c in r] for r in wb["AgentPlan"].iter_rows(min_row=2)]
    assert plan == [["AC-1.1", "boundary", "planned", "title dài 200", None, "TC-AC-1.1-aaaa01"]]


def test_write_if_changed_is_idempotent_and_repairs_a_corrupt_file(tmp_path):
    catalog = mini()
    path = tmp_path / xlsx.XLSX_PATH
    assert xlsx.write_if_changed(tmp_path, catalog) == "created" and path.is_file()
    stamp, before = path.stat().st_mtime_ns, path.read_bytes()
    assert xlsx.write_if_changed(tmp_path, copy.deepcopy(catalog)) == "unchanged" and path.read_bytes() == before and path.stat().st_mtime_ns == stamp
    changed = copy.deepcopy(catalog)
    changed["test_cases"][0]["status"] = "approved"
    assert xlsx.write_if_changed(tmp_path, changed) == "updated"
    assert xlsx.write_if_changed(tmp_path, changed) == "unchanged"
    path.write_bytes(b"not a zip")
    assert xlsx.write_if_changed(tmp_path, changed) == "updated" and xlsx.read_xlsx(path)[2] == []
    assert not list(path.parent.glob("*.tmp"))


def test_write_if_changed_rewrites_when_only_the_base_snapshot_is_stale(tmp_path):
    catalog = mini()
    xlsx.write_if_changed(tmp_path, catalog)
    qa = Edit((tmp_path / xlsx.XLSX_PATH).read_bytes()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").bytes()
    (tmp_path / xlsx.XLSX_PATH).write_bytes(qa)
    after_import = copy.deepcopy(catalog)
    after_import["test_cases"][0]["status"] = "approved"             # YAML bây giờ khớp nội dung xlsx, nhưng ảnh chụp gốc trong xlsx vẫn là bản cũ
    assert xlsx.write_if_changed(tmp_path, after_import) == "updated"
    assert xlsx.read_xlsx(tmp_path / xlsx.XLSX_PATH)[1] == xlsx.projection(after_import)


# ================= đọc an toàn =================

@pytest.mark.parametrize("blob", [b"", b"plain text SECRET", b"PK\x03\x04garbage", b"\x00" * 100])
def test_unreadable_files_are_xlsx_errors_without_content(blob):
    with pytest.raises(xlsx.XlsxError) as error:
        xlsx.read_xlsx(blob)
    assert "SECRET" not in str(error.value)


def test_missing_file_oversize_and_zip_bomb(tmp_path):
    with pytest.raises(xlsx.XlsxError):
        xlsx.read_xlsx(tmp_path / "nope.xlsx")
    with pytest.raises(xlsx.XlsxError) as big:
        xlsx.read_xlsx(b"x" * (xlsx.MAX_BYTES + 1))
    assert "lớn hơn" in str(big.value)
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 60_000_000)
    assert len(bomb.getvalue()) < xlsx.MAX_BYTES
    with pytest.raises(xlsx.XlsxError) as error:
        xlsx.read_xlsx(bomb.getvalue())
    assert "giải nén" in str(error.value)


def test_too_many_rows_and_missing_sheets_or_columns():
    data = Edit(export())
    for n in range(xlsx.MAX_ROWS + 5):
        data.wb["Uncovered"].cell(n + 2, 1).value = f"AC-{n}"
    with pytest.raises(xlsx.XlsxError) as many:
        xlsx.read_xlsx(data.bytes())
    assert "quá lớn" in str(many.value)
    with pytest.raises(xlsx.XlsxError) as sheet:
        xlsx.read_xlsx(Edit(export()).drop("Steps").bytes())
    assert "thiếu sheet Steps" in str(sheet.value)
    broken = Edit(export())
    broken.wb["TestCases"].cell(1, broken.col("TestCases", "title")).value = "tieu_de"
    with pytest.raises(xlsx.XlsxError) as col:
        xlsx.read_xlsx(broken.bytes())
    assert "thiếu cột: title" in str(col.value)


def test_formula_and_date_cells_typed_by_qa_are_reported_not_evaluated():
    qa = Edit(export())
    qa.set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "notes", "=HYPERLINK(\"http://evil.example\")")
    qa.set("TestCases", {"tc_id": "TC-AC-1.2-aaaa02"}, "notes", __import__("datetime").datetime(2026, 1, 2))
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    assert [w for w, _ in syntax] == ["TestCases!2:J", "TestCases!3:J"] or len(syntax) == 2
    assert any("công thức" in m for _, m in syntax) and any("ngày giờ" in m for _, m in syntax)
    assert "evil.example" not in json.dumps(syntax)
    result = xlsx.import_into(mini(), theirs, base, syntax)
    assert not result.ok and result.catalog["test_cases"][0].get("notes") is None


def test_columns_are_matched_by_name_so_qa_can_reorder_them():
    qa = Edit(export())
    ws = qa.wb["TestCases"]
    ws.move_range(f"A1:{chr(64 + len(xlsx.TC_COLUMNS))}{ws.max_row}", rows=0, cols=0)
    ws.insert_cols(1)                                         # đẩy mọi cột sang phải, thêm cột lạ ở đầu
    ws.cell(1, 1).value = "ghi_chu_cua_toi"
    ws.cell(2, 1).value = "bỏ qua cột lạ"
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    assert syntax == [] and xlsx.strip_rows(theirs) == xlsx.projection(mini())


# ================= import =================

def test_a_no_op_import_changes_nothing_byte_for_byte():
    catalog = mini()
    result = do_import(catalog, export(catalog))
    assert result.catalog == catalog and not (result.changed or result.added or result.removed or result.converted or result.conflicts or result.errors or result.warnings)
    assert gt_render._catalog_file(result.catalog) == gt_render._catalog_file(catalog)


def test_review_fields_are_editable_on_any_test_case_and_keep_origin():
    catalog = mini()
    qa = (Edit(export()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").set("TestCases", {"tc_id": "TC-AC-1.2-aaaa02"}, "status", "rejected")
          .set("TestCases", {"tc_id": "TC-AC-1.2-aaaa02"}, "rejected_reason", "sai PRD").set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "notes", "đã kiểm")
          .set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "priority", "low").set("Catalog", {"key": "status"}, "value", "draft"))
    result = do_import(catalog, qa.bytes())
    first, second = result.catalog["test_cases"][0], result.catalog["test_cases"][1]
    assert (first["status"], first["notes"], first["priority"], first["origin"]) == ("approved", "đã kiểm", "low", "llm") and first["technique"] == "boundary"
    assert (second["status"], second["rejected_reason"], second["origin"]) == ("rejected", "sai PRD", "llm") and result.converted == []
    assert set(result.changed) == {"TC-AC-1.1-aaaa01", "TC-AC-1.2-aaaa02"} and first["steps"] == catalog["test_cases"][0]["steps"]
    assert result.catalog["test_cases"][2] == catalog["test_cases"][2] and gt_schema.validate_catalog(xlsx._probe(result.catalog)) == []


@pytest.mark.parametrize("column,value", [("technique", "authz"), ("technique", None), ("rationale", "lý do mới"), ("preconditions", None), ("ac_refs", "AC-2.1"), ("kind", "api_contract")])
def test_any_edit_beyond_a_review_decision_converts_an_llm_case_so_regen_cannot_discard_it(column, value):
    result = do_import(mini(), Edit(export()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, column, value).bytes(), expect_ok=column != "kind")
    if column == "kind":
        assert not result.ok                                         # api_contract chỉ nhận 1 bước chỉ method+path: schema chặn
        return
    tc = result.catalog["test_cases"][0] if column != "ac_refs" else next(t for t in result.catalog["test_cases"] if t["tc_id"] == "TC-AC-1.1-aaaa01")
    assert tc["origin"] == "qa" and result.converted == ["TC-AC-1.1-aaaa01"] and protected(tc)
    assert (column not in tc) if value is None else (tc[column] == ([value] if column == "ac_refs" else value))


def test_editing_the_content_of_an_llm_case_makes_it_origin_qa_and_protects_it_from_regen():
    catalog = mini()
    qa = (Edit(export()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "title", "  tiêu đề   mới  ")
          .set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "expect_status", "201, 202").set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "query", '{"x": 1}'))
    result = do_import(catalog, qa.bytes())
    tc = result.catalog["test_cases"][0]
    assert tc["tc_id"] == "TC-AC-1.1-aaaa01" and tc["title"] == "tiêu đề mới" and tc["origin"] == "qa" and tc["status"] == "draft"
    assert tc["steps"][0]["expect"]["status"] == [201, 202] and tc["steps"][0]["request"]["query"] == {"x": 1}
    assert "sửa từ TC LLM" in tc["notes"] and result.converted == ["TC-AC-1.1-aaaa01"]
    assert protected(tc)                                                           # `regen` giữ nguyên nó
    regen = merge(result.catalog, {**catalog, "test_cases": [catalog["test_cases"][0]]})
    assert [t for t in regen.catalog["test_cases"] if t["tc_id"] == "TC-AC-1.1-aaaa01"] == [tc]


def test_assertions_are_edited_with_typed_values():
    catalog = mini()
    qa = Edit(export())
    qa.set("Assertions", {"tc_id": "TC-AC-1.1-aaaa01", "path": "$.n"}, "value", "42")
    qa.set("Assertions", {"tc_id": "TC-AC-1.1-aaaa01", "path": "$.s"}, "value_type", "number")      # "007" thành số: JSON không có 007 -> lỗi
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    assert any("number" in m or "JSON" in m for _, m in syntax)
    qa = Edit(export()).set("Assertions", {"tc_id": "TC-AC-1.1-aaaa01", "path": "$.n"}, "value", "42").set("Assertions", {"tc_id": "TC-AC-1.1-aaaa01", "path": "$.l"}, "value", "9")
    qa.add("Assertions", tc_id="TC-AC-1.1-aaaa01", step_no="1", n="99", path="$.extra", op="eq", value="true", value_type="boolean")
    qa.add("Assertions", tc_id="TC-AC-1.1-aaaa01", step_no="1", n="100", path="$.gone", op="absent")
    result = do_import(catalog, qa.bytes())
    assertions = {a["path"]: a for a in result.catalog["test_cases"][0]["steps"][0]["expect"]["json"]}
    assert assertions["$.n"]["value"] == 42 and assertions["$.l"]["value"] == 9 and assertions["$.extra"]["value"] is True and "value" not in assertions["$.gone"]
    assert result.catalog["test_cases"][0]["origin"] == "qa"


def test_new_test_cases_get_real_ids_and_origin_qa():
    catalog = mini()
    qa = Edit(export())
    for handle in ("NEW-1", "NEW-2"):
        qa.add("TestCases", tc_id=handle, ac_refs="AC-2.1, AC-1.1", title="  QA thêm   một TC ", kind="api_functional", status="")
        qa.add("Steps", tc_id=handle, step_no="1", method="get", path="/notes", expect_status="200")
        qa.add("Assertions", tc_id=handle, step_no="1", n="1", path="$", op="type", value="array")
    result = do_import(catalog, qa.bytes())
    assert len(result.added) == 2 and result.added[0] != result.added[1] and result.added[1].endswith("-2")      # nội dung giống hệt: id vẫn không trùng
    new = [t for t in result.catalog["test_cases"] if t["tc_id"] in result.added]
    assert all(t["origin"] == "qa" and t["status"] == "approved" and t["title"] == "QA thêm một TC" and t["ac_refs"] == ["AC-2.1", "AC-1.1"] for t in new)
    assert new[0]["tc_id"].startswith("TC-AC-2.1-qa-") and new[0]["steps"][0]["request"]["method"] == "GET"
    assert new[0]["steps"][0]["expect"]["json"] == [{"path": "$", "op": "type", "value": "array"}]
    assert [t["tc_id"] for t in result.catalog["test_cases"]].count(new[0]["tc_id"]) == 1
    keys = [(0, *{"AC-1.1": (0, 0), "AC-1.2": (0, 1), "AC-2.1": (1, 0)}[t["ac_refs"][0]]) for t in result.catalog["test_cases"]]
    assert keys == sorted(keys)                                                                                   # xếp theo vị trí AC đầu tiên như `merge`


def test_new_case_and_the_exported_file_agree_after_the_real_id_is_assigned():
    catalog = mini()
    qa = Edit(export())
    qa.add("TestCases", tc_id="NEW-X", ac_refs="AC-1.1", title="mới", kind="api_functional", status="approved")
    qa.add("Steps", tc_id="NEW-X", step_no="1", method="GET", path="/notes", expect_status="200")
    result = do_import(catalog, qa.bytes())
    theirs, base, syntax = xlsx.read_xlsx(xlsx.render_xlsx(result.catalog))
    assert syntax == [] and xlsx.sync_status(result.catalog, theirs, base) == "in_sync" and result.added[0] in xlsx.projection(result.catalog)["test_cases"]


@pytest.mark.parametrize("fields,steps,message", [
    (dict(ac_refs="AC-1.1", title="t", kind="api_functional", status="approved"), False, "chưa có bước nào"),
    (dict(ac_refs="AC-99.9", title="t", kind="api_functional", status="approved"), True, "ac_refs phải liệt kê"),
    (dict(ac_refs="", title="t", kind="api_functional", status="approved"), True, "ac_refs phải liệt kê"),
    (dict(ac_refs="AC-1.1", title="", kind="api_functional", status="approved"), True, "schema"),
    (dict(ac_refs="AC-1.1", title="t", kind="ui", status="approved"), True, "schema"),
    (dict(ac_refs="AC-1.1", title="t", kind="flow", status="approved"), True, "schema"),                  # flow cần từ 2 bước
])
def test_invalid_new_cases_are_errors_and_nothing_is_written(fields, steps, message):
    qa = Edit(export()).add("TestCases", tc_id="NEW-1", **fields)
    if steps:
        qa.add("Steps", tc_id="NEW-1", step_no="1", method="GET", path="/notes", expect_status="200")
    result = do_import(mini(), qa.bytes(), expect_ok=False)
    assert not result.ok and any(message in m for _, m in result.errors)


def test_bad_ids_and_orphan_rows_are_errors():
    qa = Edit(export()).add("TestCases", tc_id="weird id", ac_refs="AC-1.1", title="t", kind="api_functional", status="approved")
    qa.add("Steps", tc_id="ghost", step_no="1", method="GET", path="/notes", expect_status="200")
    qa.add("Assertions", tc_id="ghost", step_no="1", n="1", path="$", op="exists")
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    result = xlsx.import_into(mini(), theirs, base, syntax)
    assert sum("không có trong sheet TestCases" in m for _, m in result.errors) == 2 and any("tc_id mới phải là" in m for _, m in result.errors)


def test_syntax_errors_are_all_reported_with_cell_addresses_and_no_content():
    qa = Edit(export())
    qa.set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "path_params", "{not json SECRET")
    qa.set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "headers", "[1, 2]")
    qa.set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "expect_status", "ok, 20")
    qa.set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "capture", '{"v": NaN}')
    qa.set("Assertions", {"tc_id": "TC-AC-1.1-aaaa01", "path": "$.l"}, "value", "abc")
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    where = {w for w, _ in syntax}
    assert {"Steps!2:path_params", "Steps!2:headers", "Steps!2", "Steps!2:capture"} <= where and any(w.startswith("Assertions!") for w in where)
    assert "SECRET" not in json.dumps(syntax)
    result = xlsx.import_into(mini(), theirs, base, syntax)
    assert not result.ok and len(result.errors) >= 5


def test_deleting_rows_llm_is_an_error_qa_is_allowed():
    catalog = mini()
    llm = do_import(catalog, Edit(export()).delete("TestCases", tc_id="TC-AC-1.1-aaaa01").delete("Steps", tc_id="TC-AC-1.1-aaaa01").bytes(), expect_ok=False)
    assert any("đặt status rejected" in m for _, m in llm.errors) and "TC-AC-1.1-aaaa01" in [t["tc_id"] for t in llm.catalog["test_cases"]]
    qa = do_import(catalog, Edit(export()).delete("TestCases", tc_id="TC-AC-2.1-qa-bbbb03").delete("Steps", tc_id="TC-AC-2.1-qa-bbbb03").bytes())
    assert qa.removed == ["TC-AC-2.1-qa-bbbb03"] and [t["tc_id"] for t in qa.catalog["test_cases"]] == ["TC-AC-1.1-aaaa01", "TC-AC-1.2-aaaa02"]


def test_readonly_columns_cannot_be_changed():
    qa = Edit(export()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "origin", "qa").set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "evidence", "app/other.py")
    result = do_import(mini(), qa.bytes(), expect_ok=False)
    assert sorted(m for _, m in result.errors) == ["cột evidence chỉ đọc nhưng đã bị sửa (tc_id TC-AC-1.1-aaaa01)", "cột origin chỉ đọc nhưng đã bị sửa (tc_id TC-AC-1.1-aaaa01)"]


def test_waivers_uncovered_conflicts_and_catalog_status_are_editable():
    catalog = mini()
    qa = (Edit(export()).set("Waivers", {"kind": "api"}, "status", "approved").add("Waivers", kind="technique", target="POST /notes boundary:max_length:body.title@201",
          reason_code="out_of_scope", reason="  ngoài   phạm vi ", status="draft").add("Uncovered", ac_id="AC-1.2", reason="chỉ giao diện")
          .set("SpecConflicts", {"ac_id": "AC-1.2"}, "status", "resolved").set("Catalog", {"key": "status"}, "value", "approved"))
    result = do_import(catalog, qa.bytes()).catalog
    assert [(w["target"], w["status"]) for w in result["waivers"]] == [("GET /notes 500", "approved"), ("POST /notes boundary:max_length:body.title@201", "draft")]
    assert result["waivers"][1]["reason"] == "ngoài phạm vi" and result["uncovered_acs"] == [{"ac_id": "AC-1.2", "reason": "chỉ giao diện"}]
    assert result["spec_conflicts"] == [{"ac_id": "AC-1.2", "summary": "mã khác PRD", "status": "resolved", "evidence": [{"path": "app/m.py", "line": 9}]}]
    assert result["status"] == "approved"
    gone = do_import(catalog, Edit(export()).delete("Waivers", kind="api").delete("SpecConflicts", ac_id="AC-1.2").bytes()).catalog
    assert "waivers" not in gone and "spec_conflicts" not in gone
    assert do_import(catalog, Edit(export()).add("Uncovered", ac_id="AC-99", reason="x").bytes(), expect_ok=False).errors[0][0] == "Uncovered"


def test_an_xlsx_of_another_prd_is_refused():
    other = mini()
    other["prd"]["id"] = "other"
    result = do_import(mini(), export(other), expect_ok=False)
    assert result.errors == [("Catalog", "prd_id của xlsx khác catalog: đây là file của PRD khác")]


# ================= ba chiều =================

def test_three_way_merge_combines_edits_to_different_fields_and_flags_the_same_field():
    catalog = mini()
    data = export(catalog)
    yaml_side = copy.deepcopy(catalog)
    yaml_side["test_cases"][0]["notes"] = "ghi trong YAML"
    yaml_side["test_cases"][1]["status"] = "approved"
    qa = Edit(data).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").set("TestCases", {"tc_id": "TC-AC-1.2-aaaa02"}, "status", "approved")
    result = do_import(yaml_side, qa.bytes())                                      # khác trường (notes vs status) hoặc cùng giá trị: gộp được
    assert result.catalog["test_cases"][0]["status"] == "approved" and result.catalog["test_cases"][0]["notes"] == "ghi trong YAML"
    assert result.catalog["test_cases"][1]["status"] == "approved"
    clash = copy.deepcopy(catalog)
    clash["test_cases"][0]["status"] = "rejected"
    clash["test_cases"][0]["rejected_reason"] = "YAML nói loại"
    result = do_import(clash, Edit(data).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").bytes(), expect_ok=False)
    assert result.conflicts == ["TC-AC-1.1-aaaa01.status"] and result.catalog["test_cases"][0]["status"] == "rejected"   # bên YAML được giữ, không ghi đè im lặng


def test_deleting_a_qa_case_that_yaml_changed_since_export_is_a_conflict():
    catalog = mini()
    data = Edit(export(catalog)).delete("TestCases", tc_id="TC-AC-2.1-qa-bbbb03").delete("Steps", tc_id="TC-AC-2.1-qa-bbbb03").bytes()
    yaml_side = copy.deepcopy(catalog)
    yaml_side["test_cases"][2]["notes"] = "đã đổi trong YAML"
    result = do_import(yaml_side, data, expect_ok=False)
    assert result.conflicts == ["TC-AC-2.1-qa-bbbb03.<xoá>"] and len(result.catalog["test_cases"]) == 3


def test_yaml_only_changes_are_never_reverted_by_a_stale_xlsx():
    catalog = mini()
    data = export(catalog)
    newer = copy.deepcopy(catalog)
    newer["test_cases"][0]["status"] = "approved"
    newer["test_cases"].pop()                                                      # YAML xoá một TC kể từ lúc xuất
    result = do_import(newer, data)
    assert result.catalog == newer and not result.changed and not result.added
    theirs, base, _ = xlsx.read_xlsx(data)
    assert xlsx.sync_status(newer, theirs, base) == "yaml_ahead"


def test_without_the_meta_sheet_the_import_is_two_way_with_a_warning():
    catalog = mini()
    qa = Edit(export(catalog)).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").drop("_meta")
    yaml_side = copy.deepcopy(catalog)
    yaml_side["test_cases"][0]["status"] = "rejected"
    yaml_side["test_cases"][0]["rejected_reason"] = "x"
    theirs, base, syntax = xlsx.read_xlsx(qa.bytes())
    assert base is None
    result = xlsx.import_into(yaml_side, theirs, base, syntax)
    assert result.ok and result.catalog["test_cases"][0]["status"] == "approved" and any("hai chiều" in m for _, m in result.warnings)
    assert xlsx.sync_status(catalog, theirs, None) == "diverged"


@pytest.mark.parametrize("damage", ["garbage", "wrong_version", "bad_base64"])
def test_a_damaged_meta_sheet_is_treated_as_missing(damage):
    qa = Edit(export())
    ws = qa.wb["_meta"]
    if damage == "wrong_version":
        ws["B1"].value = "99"
    else:
        for row in ws.iter_rows(min_row=5):
            if row[0].value == "base":
                row[1].value = "!!!not-base64!!!" if damage == "bad_base64" else "AAAA"
    _, base, _ = xlsx.read_xlsx(qa.bytes())
    assert base is None


def test_sync_status_matrix():
    catalog = mini()
    theirs, base, _ = xlsx.read_xlsx(export(catalog))
    assert xlsx.sync_status(catalog, theirs, base) == "in_sync"
    edited = copy.deepcopy(catalog)
    edited["test_cases"][0]["notes"] = "n"
    assert xlsx.sync_status(edited, theirs, base) == "yaml_ahead"
    qa_theirs, qa_base, _ = xlsx.read_xlsx(Edit(export(catalog)).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "notes", "x").bytes())
    assert xlsx.sync_status(catalog, qa_theirs, qa_base) == "xlsx_ahead" and xlsx.sync_status(edited, qa_theirs, qa_base) == "diverged"


# ================= cổng validate =================

def build_repo(tmp_path, catalog=None, facts=None) -> Path:
    root = tmp_path / "repo"
    catalog = catalog or mini()
    gt_dir = root / GT
    (gt_dir / "tests_gt").mkdir(parents=True)
    for rendered in gt_render.render(catalog, sut_root=root):
        if rendered.path.startswith(".qc-agent/ground-truth/"):
            target = root / rendered.path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(rendered.content, encoding="utf-8", newline="\n")
    (gt_dir / "module-map.yaml").write_text("version: 1\nstatus: approved\nmodules: []\n", encoding="utf-8")
    xlsx.write_if_changed(root, catalog, facts)
    return root


def findings(root):
    result = gt_check.check(root)
    return [m for _, m in result.errors], [m for _, m in result.warnings]


def test_validate_requires_xlsx_to_match_yaml(tmp_path):
    root = build_repo(tmp_path)
    errors, warnings = findings(root)
    assert not [e for e in errors if "xlsx" in e or "import-xlsx" in e or "export-xlsx" in e]
    path = root / xlsx.XLSX_PATH
    qa = Edit(path.read_bytes()).set("TestCases", {"tc_id": "TC-AC-1.1-aaaa01"}, "status", "approved").bytes()
    path.write_bytes(qa)
    errors, _ = findings(root)
    assert any("chạy `qc-agent gt import-xlsx`" in e for e in errors)                 # quyết định của QA trong xlsx mà gate không thấy: LỖI
    both = load_catalog(root)
    both["test_cases"][1]["notes"] = "YAML cũng đổi"
    save_catalog(root, both)
    errors, _ = findings(root)
    assert any("khác nhau" in e and "import-xlsx" in e for e in errors)
    path.write_bytes(xlsx.render_xlsx(mini()))                                          # xlsx quay về bản gốc; YAML vẫn mới hơn
    errors, warnings = findings(root)
    assert not [e for e in errors if "xlsx" in e.lower()] and any("Excel đã cũ" in w and "export-xlsx" in w for w in warnings)


def test_validate_reports_syntax_problems_and_unreadable_xlsx(tmp_path):
    root = build_repo(tmp_path)
    path = root / xlsx.XLSX_PATH
    path.write_bytes(Edit(path.read_bytes()).set("Steps", {"tc_id": "TC-AC-1.1-aaaa01"}, "path_params", "{oops").bytes())
    errors, _ = findings(root)
    assert any("Steps!2:path_params: không phải JSON hợp lệ" in e for e in errors)
    path.write_bytes(b"corrupt SECRET")
    with pytest.raises(gt_check.GTCheckError) as error:
        gt_check.check(root)
    assert "SECRET" not in str(error.value) and "test-cases.xlsx" in str(error.value)


def test_without_an_xlsx_the_gate_and_core_never_import_openpyxl(tmp_path):
    root = build_repo(tmp_path)
    (root / xlsx.XLSX_PATH).unlink()
    code = ("import sys\nfrom qc_agent.groundtruth import check\nimport qc_agent.core.cli, qc_agent.core.engine\n"
            f"r = check.check({str(root)!r})\nassert not any(m for _, m in r.errors if 'xlsx' in m)\n"
            "bad = [m for m in ('openpyxl', 'defusedxml', 'qc_agent.groundtruth.xlsx') if m in sys.modules]\nassert not bad, bad\nprint('ok')")
    done = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=Path(__file__).resolve().parent.parent)
    assert done.returncode == 0 and done.stdout.strip() == "ok", done.stdout + done.stderr


# ================= CLI =================

def sheet_path(sut):
    return sut / xlsx.XLSX_PATH


def test_generate_writes_the_xlsx_and_it_can_be_turned_off(tmp_path, capsys, fake, monkeypatch):
    sut = generated(tmp_path, capsys)
    theirs, base, syntax = xlsx.read_xlsx(sheet_path(sut))
    assert syntax == [] and base == xlsx.projection(load_catalog(sut)) and xlsx.strip_rows(theirs) == base
    wb = load_workbook(sheet_path(sut))
    assert any(r[0].value == "api" and r[2].value == "gap" for r in wb["Coverage"].iter_rows(min_row=2))   # generate có OpenAPI nên Coverage có cả technique/API
    off = make_sut(tmp_path, "off")
    assert gt(capsys, *generate_args(off, "--no-xlsx", egress=tmp_path / "e1"))[0] == 0 and not sheet_path(off).exists()
    monkeypatch.setenv("QC_GT_XLSX", "false")
    env_off = make_sut(tmp_path, "envoff")
    assert gt(capsys, *generate_args(env_off, egress=tmp_path / "e2"))[0] == 0 and not sheet_path(env_off).exists()


def qa_edit(sut, fn):
    sheet_path(sut).write_bytes(fn(Edit(sheet_path(sut).read_bytes())).bytes())


def test_import_xlsx_cli_end_to_end(tmp_path, capsys, fake):
    sut = generated(tmp_path, capsys)
    first = load_catalog(sut)["test_cases"][0]["tc_id"]
    qa_edit(sut, lambda e: e.set("TestCases", {"tc_id": first}, "status", "approved").set("TestCases", {"tc_id": first}, "notes", "QA ok")
            .add("TestCases", tc_id="NEW-1", ac_refs="AC-1.1", title="QA thêm", kind="api_functional", status="approved")
            .add("Steps", tc_id="NEW-1", step_no="1", method="GET", path="/notes", expect_status="200"))
    code, out, err = gt(capsys, "validate", "--sut-root", sut)
    assert code == 1 and "chạy `qc-agent gt import-xlsx`" in out
    before = (sut / GT / "test-cases.yaml").read_bytes()
    code, out, err = gt(capsys, "import-xlsx", "--sut-root", sut, "--dry-run")
    assert code == 0 and "1 TC mới" in out and "--dry-run: chưa ghi gì" in out and (sut / GT / "test-cases.yaml").read_bytes() == before
    code, out, err = gt(capsys, "import-xlsx", "--sut-root", sut)
    assert code == 0, out + err
    catalog = load_catalog(sut)
    mine = next(t for t in catalog["test_cases"] if t["tc_id"] == first)
    new = next(t for t in catalog["test_cases"] if t["origin"] == "qa")
    assert (mine["status"], mine["notes"]) == ("approved", "QA ok") and new["tc_id"].startswith("TC-AC-1.1-qa-") and new["status"] == "approved"
    assert "mới" in out and new["tc_id"] in out and "created" not in out
    theirs, base, _ = xlsx.read_xlsx(sheet_path(sut))
    assert new["tc_id"] in base["test_cases"] and "NEW-1" not in base["test_cases"]                    # xlsx được xuất lại với id thật và ảnh chụp gốc mới
    code, out, _ = gt(capsys, "validate", "--sut-root", sut)
    assert "import-xlsx" not in out and "export-xlsx" not in out                                         # không còn lệch (còn lỗi khác như draft là chuyện khác)
    assert gt(capsys, "import-xlsx", "--sut-root", sut)[0] == 0                                          # chạy lại: không thay gì


def test_import_xlsx_cli_conflicts_errors_and_unreadable(tmp_path, capsys, fake):
    sut = generated(tmp_path, capsys)
    ids = [t["tc_id"] for t in load_catalog(sut)["test_cases"]]
    catalog = load_catalog(sut)
    catalog["test_cases"][0]["status"] = "rejected"
    catalog["test_cases"][0]["rejected_reason"] = "YAML loại"
    save_catalog(sut, catalog)
    qa_edit(sut, lambda e: e.set("TestCases", {"tc_id": ids[0]}, "status", "approved"))
    before = (sut / GT / "test-cases.yaml").read_bytes()
    code, out, err = gt(capsys, "import-xlsx", "--sut-root", sut)
    assert code == 1 and f"XUNG ĐỘT  {ids[0]}.status" in out and "KHÔNG ghi gì" in out and (sut / GT / "test-cases.yaml").read_bytes() == before
    sheet_path(sut).write_bytes(b"junk SECRET")
    code, out, err = gt(capsys, "import-xlsx", "--sut-root", sut)
    assert code == 3 and "SECRET" not in out + err and "LỖI" in err
    assert gt(capsys, "import-xlsx", "--sut-root", sut, "--xlsx", tmp_path / "missing.xlsx")[0] == 3


def test_export_xlsx_cli_refuses_to_lose_unimported_work(tmp_path, capsys, fake):
    sut = generated(tmp_path, capsys)
    first = load_catalog(sut)["test_cases"][0]["tc_id"]
    assert "unchanged" in gt(capsys, "export-xlsx", "--sut-root", sut)[1]
    qa_edit(sut, lambda e: e.set("TestCases", {"tc_id": first}, "status", "approved"))
    code, out, err = gt(capsys, "export-xlsx", "--sut-root", sut)
    assert code == 1 and "chưa import" in err and xlsx.read_xlsx(sheet_path(sut))[0]["test_cases"][first]["status"] == "approved"
    code, out, err = gt(capsys, "export-xlsx", "--sut-root", sut, "--force")
    assert code == 0 and "updated" in out and xlsx.read_xlsx(sheet_path(sut))[0]["test_cases"][first]["status"] == "draft"        # --force bỏ sửa chưa import
    catalog = load_catalog(sut)
    catalog["test_cases"][0]["notes"] = "sửa tay YAML"
    save_catalog(sut, catalog)
    assert gt(capsys, "export-xlsx", "--sut-root", sut)[0] == 0 and xlsx.read_xlsx(sheet_path(sut))[0]["test_cases"][first]["notes"] == "sửa tay YAML"
    sheet_path(sut).write_bytes(b"corrupt")
    assert gt(capsys, "export-xlsx", "--sut-root", sut)[0] == 0 and xlsx.read_xlsx(sheet_path(sut))[2] == []                            # file hỏng thì xuất lại được


def test_regen_absorbs_unimported_xlsx_edits_before_merging(tmp_path, capsys, fake):
    sut = generated(tmp_path, capsys)
    first = load_catalog(sut)["test_cases"][0]["tc_id"]
    qa_edit(sut, lambda e: e.set("TestCases", {"tc_id": first}, "status", "approved").set("TestCases", {"tc_id": first}, "notes", "chưa import"))
    code, out, err = gt(capsys, "regen", *generate_args(sut, egress=tmp_path / "e2")[1:])
    assert code == 0, out + err
    kept = next(t for t in load_catalog(sut)["test_cases"] if t["tc_id"] == first)
    assert (kept["status"], kept["notes"]) == ("approved", "chưa import")                       # không bị regen ghi đè
    assert "đã gộp" in out and "từ .qc-agent/ground-truth/test-cases.xlsx trước khi regen" in out
    theirs, base, _ = xlsx.read_xlsx(sheet_path(sut))
    assert xlsx.sync_status(load_catalog(sut), theirs, base) == "in_sync"                       # regen xuất lại xlsx cho khớp


def test_regen_stops_on_a_conflict_instead_of_overwriting(tmp_path, capsys, fake):
    sut = generated(tmp_path, capsys)
    ids = [t["tc_id"] for t in load_catalog(sut)["test_cases"]]
    catalog = load_catalog(sut)
    catalog["test_cases"][0].update(status="rejected", rejected_reason="YAML")
    save_catalog(sut, catalog)
    qa_edit(sut, lambda e: e.set("TestCases", {"tc_id": ids[0]}, "status", "approved"))
    before = (sut / GT / "test-cases.yaml").read_bytes()
    code, out, err = gt(capsys, "regen", *generate_args(sut, egress=tmp_path / "e2")[1:])
    assert code == 3 and "không gộp được vào YAML" in err and f"xung đột {ids[0]}.status" in err and (sut / GT / "test-cases.yaml").read_bytes() == before
    assert fake.count == 1                                                                      # dừng TRƯỚC khi gọi LLM lần hai


def test_an_xml_entity_expansion_payload_is_refused_not_expanded():
    """openpyxl dùng defusedxml khi đã cài: một worksheet chứa DOCTYPE/entity phải bị từ chối (XlsxError), không bị bung ra."""
    src = zipfile.ZipFile(io.BytesIO(export()))
    bomb = ('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&b;</t></is></c></row></sheetData></worksheet>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name in src.namelist():
            archive.writestr(name, bomb if name == "xl/worksheets/sheet2.xml" else src.read(name))
    with pytest.raises(xlsx.XlsxError) as error:
        xlsx.read_xlsx(out.getvalue())
    assert "aaaa" not in str(error.value)
